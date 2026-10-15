"""Regression coverage for interrupted imports and bounded embedding workers."""
from __future__ import annotations

import asyncio
import io
import json
import time

import numpy as np
import pytest
from openpyxl import Workbook

from app import repo
from app.catalog.index import load_index
from app.catalog.texts import vector_texts
from app.db import bump_version, create_pool, migrate
from app.gemini.fake import FakeGemini
from app.gemini.pool import GeminiError
from app.ingest import importer
from app.ingest.vectors import RateLimiter, embed_demo
from app.jobs import JobRunner
from tests.conftest import make_settings


async def catalog(db_url, n=4):
    pool = await create_pool(db_url)
    await migrate(pool)
    demo = await repo.create_demo(pool, {"slug": "checkpoint", "name": "Test", "embed_dim": 64})
    async with pool.acquire() as conn, conn.transaction():
        await repo.upsert_products(conn, demo["id"], [
            {"sku": str(i), "name": f"Sản phẩm {i}", "short_description": f"Nhu cầu {i}"} for i in range(n)])
        await bump_version(conn, demo["id"])
    return pool, demo["id"]


async def run_embed(pool, gem, did, **kwargs):
    return await embed_demo(pool, gem, did, per_minute=10000, timeout_ms=2000, batch=1, **kwargs)


async def test_failed_batch_keeps_checkpoint_and_retry_only_embeds_missing(db_url):
    pool, did = await catalog(db_url)

    class FailSecond(FakeGemini):
        async def embed(self, *args, **kwargs):
            if self.calls.get("embed_doc", 0) == 1:
                raise GeminiError("quota_exhausted", "test")
            return await super().embed(*args, **kwargs)

    try:
        gem = FailSecond(make_settings(db_url))
        with pytest.raises(GeminiError, match="quota_exhausted"):
            await run_embed(pool, gem, did, concurrency=1)
        saved = len(await repo.existing_hashes(pool, did))
        assert saved == 1
        gem = FakeGemini(make_settings(db_url))
        result = await run_embed(pool, gem, did, concurrency=2)
        assert result["reused"] == saved
        assert result["total_vectors"] == result["embedded"] + saved == 8
        assert gem.calls["embed_doc"] == 7
        assert (await load_index(pool, "checkpoint")).n_vectors == 8
    finally:
        await pool.close()


async def test_shutdown_resumes_same_durable_job_and_reuses_batches(db_url):
    pool, did = await catalog(db_url)
    blocked = asyncio.Event()

    class PausingGemini(FakeGemini):
        async def embed(self, *args, **kwargs):
            if self.calls.get("embed_doc", 0) == 1:
                blocked.set()
                await asyncio.Event().wait()
            return await super().embed(*args, **kwargs)

    first = JobRunner(pool)
    gem = PausingGemini(make_settings(db_url))

    async def handler(ctx, demo_id, params):
        return await run_embed(pool, gem, demo_id, concurrency=1, checkpoint=ctx.checkpoint)

    first.register("embed", handler)
    try:
        jid = await first.submit("embed", did, {"requested_version": 2})
        await asyncio.wait_for(blocked.wait(), 3)
        await first.cancel_all()
        assert (await repo.get_job(pool, jid))["status"] == "queued"
        assert len(await repo.existing_hashes(pool, did)) == 1
        second = JobRunner(pool)
        gem = FakeGemini(make_settings(db_url))
        second.register("embed", handler)
        await second.recover()
        await second.wait(jid)
        job = await repo.get_job(pool, jid)
        assert job["status"] == "done", job
        assert job["result"]["reused"] == 1
        assert len(await repo.existing_hashes(pool, did)) == 8
    finally:
        await first.cancel_all()
        await pool.close()


async def test_workers_are_parallel_but_bounded_and_identical_texts_are_shared(db_url):
    pool, did = await catalog(db_url)

    class TrackingGemini(FakeGemini):
        active = peak = 0

        async def embed(self, *args, **kwargs):
            self.active += 1
            self.peak = max(self.peak, self.active)
            try:
                await asyncio.sleep(.02)
                return await super().embed(*args, **kwargs)
            finally:
                self.active -= 1

    gem = TrackingGemini(make_settings(db_url, embed_concurrency=2))
    try:
        async with pool.acquire() as conn, conn.transaction():
            await conn.execute("UPDATE products SET short_description='same profile' WHERE demo_id=$1", did)
            await bump_version(conn, did)
        result = await run_embed(pool, gem, did, concurrency=2)
        assert gem.peak == 2
        assert result["total_vectors"] == 8
        assert result["api_documents"] == 5
    finally:
        await pool.close()


async def test_edit_during_embedding_rejects_old_snapshot(db_url):
    pool, did = await catalog(db_url)

    class EditingGemini(FakeGemini):
        async def embed(self, *args, **kwargs):
            if not self.calls:
                async with pool.acquire() as conn, conn.transaction():
                    await repo.upsert_products(conn, did, [{"sku": "0", "name": "Changed"}])
                    await bump_version(conn, did)
            return await super().embed(*args, **kwargs)

    try:
        with pytest.raises(GeminiError, match="catalog_changed"):
            await run_embed(pool, EditingGemini(make_settings(db_url)), did, concurrency=1)
        assert not await repo.existing_hashes(pool, did)
        assert (await load_index(pool, "checkpoint")).n_vectors == 0
        result = await run_embed(pool, FakeGemini(make_settings(db_url)), did, concurrency=1)
        assert result["total_vectors"] == 7  # blank profile is not sent
    finally:
        await pool.close()


@pytest.mark.parametrize("bad", ["shape", "nan", "zero"])
async def test_invalid_embedding_does_not_commit_batch(db_url, bad):
    pool, did = await catalog(db_url)

    class InvalidGemini(FakeGemini):
        async def embed(self, *args, **kwargs):
            if bad == "shape":
                return np.ones((1, 3))
            return np.full((1, 64), np.nan if bad == "nan" else 0)

    try:
        with pytest.raises(GeminiError, match="partial_failure") as failed:
            await run_embed(pool, InvalidGemini(make_settings(db_url)), did, concurrency=1)
        assert failed.value.result["failed_vectors"] == 8
        assert all(item["reason"] == "bad_output" for item in failed.value.result["dead_letters"])
        assert not await repo.existing_hashes(pool, did)
    finally:
        await pool.close()


async def test_jobs_deduplicate_and_serialize_catalog_without_starving_other_catalogs(db_url):
    pool, did = await catalog(db_url)
    other = (await repo.create_demo(pool, {"slug": "other", "name": "Other"}))["id"]
    runner = JobRunner(pool, max_concurrency=2)
    started, release = asyncio.Event(), asyncio.Event()
    calls = []

    async def handler(ctx, demo_id, params):
        calls.append((demo_id, params["requested_version"]))
        if demo_id == did and params["requested_version"] == 1:
            started.set()
            await release.wait()
        return {}

    runner.register("embed", handler)
    try:
        first = await runner.submit("embed", did, {"requested_version": 1})
        assert await runner.submit("embed", did, {"requested_version": 1}) == first
        await started.wait()
        queued = await runner.submit("embed", did, {"requested_version": 2})
        unrelated = await runner.submit("embed", other, {"requested_version": 1})
        await asyncio.wait_for(runner.wait(unrelated), 3)
        assert (did, 2) not in calls
        assert (await repo.get_job(pool, queued))["status"] == "queued"
        release.set()
        await runner.wait(first)
        await runner.wait(queued)
        assert calls.count((did, 2)) == 1
    finally:
        await runner.cancel_all()
        await pool.close()


async def test_rate_limiter_reserves_documents_across_concurrent_workers():
    limiter = RateLimiter(3, window_s=.06)
    started = time.monotonic()
    await asyncio.gather(limiter.acquire(2), limiter.acquire(2))
    assert time.monotonic() - started >= .055
    with pytest.raises(ValueError):
        await limiter.acquire(4)


def workbook(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("content,filename,data", [
    (b"not an Excel file", "broken.xlsx", {}),
    (b"", "empty.csv", {}),
    (b"name\nA", "file.exe", {}),
    (b"name\nA", "file.csv", {"mapping": "[]"}),
    (b"name\nA", "file.csv", {"mapping": '{"name":"missing"}'}),
    (b"name\nA", "file.csv", {"definitions": '[{"key":"BAD!","label":"x"}]'}),
])
def test_invalid_imports_return_validation_errors_without_creating_table(client, content, filename, data):
    response = client.post("/v1/tables", files={"file": (filename, content)}, data={"name": "Bad", **data})
    assert response.status_code == 422, response.text
    assert client.get("/v1/tables").json() == []


def test_failed_database_import_rolls_back_table_creation(client, monkeypatch):
    async def fail(*args, **kwargs):
        raise RuntimeError("simulated storage failure")

    monkeypatch.setattr(repo, "upsert_products", fail)
    with pytest.raises(RuntimeError, match="simulated storage failure"):
        client.post("/v1/tables", files={"file": ("data.xlsx", workbook([["name"], ["A"]]))}, data={"name": "Rollback"})
    assert client.get("/v1/tables").json() == []


async def test_parser_runs_outside_event_loop_and_enforces_row_limit(monkeypatch):
    original = importer.read_table
    started = asyncio.Event()
    loop = asyncio.get_running_loop()

    def slow_read(*args):
        loop.call_soon_threadsafe(started.set)
        time.sleep(.15)
        return original(*args)

    monkeypatch.setattr(importer, "read_table", slow_read)
    task = asyncio.create_task(importer.plan_import(None, None, b"name\nA", "a.csv"))
    await asyncio.wait_for(started.wait(), 1)
    await asyncio.sleep(.01)
    assert not task.done()  # the event loop stayed responsive during the parse
    assert (await task)["n_rows"] == 1
    monkeypatch.setattr(importer, "MAX_ROWS", 1)
    with pytest.raises(ValueError, match="10.000"):
        await importer.plan_import(None, None, b"name\nA\nB", "a.csv")


def test_vector_documents_skip_empty_profiles_and_preserve_existing_documents():
    long_name = "Tên sản phẩm " * 500
    docs = vector_texts({"name": long_name}, [])
    assert docs == [("name", long_name)]


def test_excel_dates_and_nonfinite_numbers_are_json_safe():
    from datetime import datetime

    row = importer.json_safe_row({"date": datetime(2026, 10, 15), "nan": float("nan"), "inf": float("inf")})
    assert json.loads(json.dumps(row, allow_nan=False)) == {"date": "2026-10-15T00:00:00", "nan": None, "inf": None}


@pytest.mark.parametrize("code,tries", [(503, 3), (400, 1), (429, 3)])
async def test_provider_retries_transient_failures_but_not_invalid_requests(db_url, monkeypatch, code, tries):
    from google.genai.errors import APIError
    from app.gemini import pool as gem_pool

    sleeps = []

    async def no_sleep(seconds):
        sleeps.append(seconds)

    monkeypatch.setattr(gem_pool.asyncio, "sleep", no_sleep)
    gem = gem_pool.GeminiPool(make_settings(db_url, key_cooldown_s=0))
    gem.clients = [object()]
    gem.cooldown = [0.0]
    attempts = 0

    async def call(_client):
        nonlocal attempts
        attempts += 1
        if attempts < 3:
            raise APIError(code, {"error": {"message": "test", "code": code}})
        return "ok"

    # Keep time-based cooldown out of this deterministic retry test.
    monkeypatch.setattr(gem, "_pick", lambda excluded: 0)
    if code == 400:
        with pytest.raises(GeminiError, match="400"):
            await gem._call("embed_doc", 1, "fake", 1000, call, max_key_tries=4)
    else:
        result, _ = await gem._call("embed_doc", 1, "fake", 1000, call, max_key_tries=4)
        assert result == "ok" and len(sleeps) == 2
    assert attempts == tries


def test_retry_failed_embedding_preserves_imported_rows(client):
    import time

    client.app.state.gemini.fail["embed_doc"] = "quota_exhausted"
    response = client.post("/v1/tables", files={"file": ("data.xlsx", workbook([["name"], ["A"]]))},
                           data={"name": "Retry", "slug": "retry"})
    assert response.status_code == 201
    jid = response.json()["embed_job"]
    for _ in range(100):
        failed = client.get(f"/v1/admin/jobs/{jid}").json()
        if failed["status"] == "error":
            break
        time.sleep(.01)
    assert failed["status"] == "error"
    assert client.get("/v1/tables/retry/rows").json()["total"] == 1
    client.app.state.gemini.fail.clear()
    response = client.post(f"/v1/admin/jobs/{jid}/retry")
    assert response.status_code == 200
    from tests.conftest import wait_job

    completed = wait_job(client, response.json()["job_id"])
    assert completed["result"]["total_vectors"] == 1
    assert client.post(f"/v1/admin/jobs/{completed['id']}/retry").status_code == 409


async def test_bad_document_is_isolated_and_dead_letters_are_durable(db_url):
    pool, did = await catalog(db_url)
    bad_text = "Sản phẩm 2"

    class PoisonedDocument(FakeGemini):
        async def embed(self, texts, *args, **kwargs):
            if bad_text in texts:
                raise GeminiError("invalid_input", "document rejected")
            return await super().embed(texts, *args, **kwargs)

    gem = PoisonedDocument(make_settings(db_url))
    runner = JobRunner(pool)

    async def handler(ctx, demo_id, params):
        return await embed_demo(pool, gem, demo_id, per_minute=10000, timeout_ms=2000,
                                batch=8, concurrency=1, checkpoint=ctx.checkpoint)

    runner.register("embed", handler)
    try:
        jid = await runner.submit("embed", did, {})
        await runner.wait(jid)
        job = await repo.get_job(pool, jid)
        assert job["status"] == "error"
        assert job["result"]["failed_vectors"] == 1
        assert job["result"]["completed_vectors"] == 7
        assert len(await repo.existing_hashes(pool, did)) == 7
        assert bad_text not in json.dumps(job["result"])  # no document contents in diagnostics
        assert job["result"]["dead_letters"][0]["reason"] == "invalid_input"
        # The next run reuses all seven good vectors, only requesting the failed document.
        healthy = FakeGemini(make_settings(db_url))
        result = await run_embed(pool, healthy, did, concurrency=1)
        assert result["reused"] == 7 and result["embedded"] == 1
        assert len(await repo.existing_hashes(pool, did)) == 8
    finally:
        await runner.cancel_all()
        await pool.close()


async def test_offline_timeout_retries_are_bounded(db_url, monkeypatch):
    from app.gemini import pool as gem_pool

    async def no_sleep(_seconds):
        pass

    monkeypatch.setattr(gem_pool.asyncio, "sleep", no_sleep)
    gem = gem_pool.GeminiPool(make_settings(db_url))
    gem.clients, gem.cooldown = [object()], [0.0]
    calls = 0

    async def timeout(_client):
        nonlocal calls
        calls += 1
        raise asyncio.TimeoutError()

    with pytest.raises(GeminiError, match="timeout"):
        await gem._call("embed_doc", 1, "fake", 10, timeout, max_key_tries=3)
    assert calls == 3
    assert gem.usage.summary()["embed_doc"]["error_codes"]["timeout"] == 3


async def test_database_transient_failure_rolls_back_and_retries_idempotently(db_url, monkeypatch):
    import asyncpg

    pool, did = await catalog(db_url)
    original = repo.add_vectors
    failed_once = False

    async def transient(conn, demo_id, rows):
        nonlocal failed_once
        if not failed_once:
            failed_once = True
            raise asyncpg.ConnectionDoesNotExistError("simulated disconnected database")
        return await original(conn, demo_id, rows)

    monkeypatch.setattr(repo, "add_vectors", transient)
    try:
        result = await run_embed(pool, FakeGemini(make_settings(db_url)), did, concurrency=1)
        assert failed_once and result["embedded"] == 8
        assert len(await repo.existing_hashes(pool, did)) == 8
        assert (await repo.get_demo_by_id(pool, did))["data_version"] == 11
    finally:
        await pool.close()
