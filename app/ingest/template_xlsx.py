"""Sinh template.xlsx theo definitions của demo: sheet products (cột lõi + thuộc tính), sheet definitions,
data validation cho enum/bool."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CORE = ["sku", "name", "category", "brand", "price", "price_unit", "in_stock", "short_description",
        "description", "aliases", "use_cases", "sample_questions"]
HINTS = {"aliases": "ngăn cách bởi |", "use_cases": "ngăn cách bởi |", "sample_questions": "ngăn cách bởi |",
         "in_stock": "có/không", "price": "VND, ví dụ 9890000 hoặc 9,89tr"}


def build_template(defs: list[dict], products: list[dict] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    headers = CORE + [f"{d['label']} ({d['unit']})" if d.get("unit") else d["label"] for d in defs]
    ws.append(headers)
    ws.append([HINTS.get(h, "") for h in CORE] + [d["type"] for d in defs])
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = bold
    for c in ws[2]:
        c.font = Font(italic=True, color="888888")
        c.fill = fill
    for p in products or []:
        attrs = p.get("attributes") or {}
        row = [p.get(k) if k not in ("aliases", "use_cases", "sample_questions") else " | ".join(p.get(k) or [])
               for k in CORE]
        for d in defs:
            v = attrs.get(d["key"])
            row.append(", ".join(v) if isinstance(v, list) else ("có" if v is True else "không" if v is False else v))
        ws.append(row)
    last = max(ws.max_row, 500)
    for j, d in enumerate(defs):
        col = get_column_letter(len(CORE) + j + 1)
        if d["type"] == "enum":
            vals = [ev["value"] if isinstance(ev, dict) else ev for ev in d.get("enum_values") or []]
            src = ",".join(vals)
            if vals and len(src) < 250:
                dv = DataValidation(type="list", formula1=f'"{src}"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col}3:{col}{last}")
        elif d["type"] == "bool":
            dv = DataValidation(type="list", formula1='"có,không"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}3:{col}{last}")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    wd = wb.create_sheet("definitions")
    cols = ["key", "label", "type", "unit", "unit_aliases", "label_aliases", "enum_values", "filterable",
            "searchable", "weight"]
    wd.append(cols)
    for c in wd[1]:
        c.font = bold
    for d in defs:
        wd.append([
            d["key"], d["label"], d["type"], d.get("unit"), " | ".join(d.get("unit_aliases") or []),
            " | ".join(d.get("label_aliases") or []),
            " | ".join((ev["value"] + (f" ({', '.join(ev.get('aliases') or [])})" if ev.get("aliases") else ""))
                       if isinstance(ev, dict) else str(ev) for ev in d.get("enum_values") or []),
            d.get("filterable"), d.get("searchable"), d.get("weight"),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
